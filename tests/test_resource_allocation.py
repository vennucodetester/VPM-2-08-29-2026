import unittest
import os
import tempfile
from openpyxl import load_workbook

from models.task_node import TaskNode
from utils.resource_allocation import (
    ResourceDefinition, accept_conflict, analyze_projects,
    next_available_start,
)
from utils.vpmt_io import load_projects, save_projects
from utils.excel_export import export_projects


def task(name, start, end, tokens, parent=None, status="Not Started"):
    node = TaskNode(name, parent=parent)
    node.start_date, node.end_date = start, end
    node.status = status
    node.task_tokens = [dict(value) for value in tokens]
    return node


def token(resource_id, label, kind="case"):
    return {"id": resource_id, "label": label, "kind": kind,
            "header": kind.title()}


def project(project_id, roots, holidays=(), weekends=True, name=None):
    return {"id": project_id, "name": name or project_id, "roots": roots,
            "metadata": {"holidays": list(holidays),
                         "exclude_weekends": weekends}}


class ResourceAllocationTests(unittest.TestCase):
    def analyze_pair(self, first, second, definitions=(), resolutions=()):
        return analyze_projects([project("P", [first, second])], definitions,
                                resolutions)

    def test_nonoverlap_different_id_and_same_label_rules(self):
        first = task("A", "2026-09-01", "2026-09-03", [token("case-1", "C")])
        later = task("B", "2026-09-04", "2026-09-07", [token("case-1", "C")])
        self.assertFalse(self.analyze_pair(first, later).conflicts)
        other = task("C", "2026-09-01", "2026-09-03", [token("case-2", "C")])
        self.assertFalse(self.analyze_pair(first, other).conflicts)

    def test_shared_boundary_partial_and_identical_spans(self):
        first = task("A", "2026-09-01", "2026-09-10", [token("case-1", "C")])
        boundary = task("B", "2026-09-10", "2026-09-11", [token("case-1", "C")])
        result = self.analyze_pair(first, boundary)
        self.assertEqual(1, result.conflicts[0].overlap_workdays)
        self.assertEqual("2026-09-10", result.conflicts[0].overlap_start)

        partial = task("B", "2026-09-07", "2026-09-14", [token("case-1", "C")])
        result = self.analyze_pair(first, partial)
        self.assertEqual(("2026-09-07", "2026-09-10", 4),
                         (result.conflicts[0].overlap_start,
                          result.conflicts[0].overlap_end,
                          result.conflicts[0].overlap_workdays))
        identical = task("B", "2026-09-01", "2026-09-10", [token("case-1", "C")])
        self.assertEqual(8, self.analyze_pair(first, identical).conflicts[0].overlap_workdays)

    def test_rename_same_id_and_different_rooms(self):
        first = task("A", "2026-09-01", "2026-09-03", [
            token("case-1", "Old"), token("room-1", "Room 1", "room")])
        second = task("B", "2026-09-01", "2026-09-03", [
            token("case-1", "Renamed"), token("room-3", "Room 3", "room")])
        result = self.analyze_pair(first, second)
        case_conflict = next(c for c in result.conflicts if c.resource_id == "case-1")
        self.assertEqual(("Room 1", "Room 3"), case_conflict.locations)
        self.assertIn("Room 3", first.schedule_conflicts[0])

    def test_missing_room_and_multiple_cases_on_one_task(self):
        first = task("A", "2026-09-01", "2026-09-02", [
            token("case-1", "C1"), token("case-2", "C2")])
        second = task("B", "2026-09-01", "2026-09-02", [
            token("case-1", "C1"), token("case-2", "C2")])
        result = self.analyze_pair(first, second)
        self.assertEqual({"case-1", "case-2"},
                         {c.resource_id for c in result.conflicts})
        self.assertEqual(("Location not assigned",), result.conflicts[0].locations)

    def test_rooms_activities_and_custom_metadata_are_generic_resources(self):
        tokens = [
            token("room-1", "Room 1", "room"),
            token("doe", "DOE", "activity"),
            token("fixture-1", "Fixture A", "custom_fixture"),
            token("phase-1", "Design", "phase"),
        ]
        first = task("A", "2026-09-01", "2026-09-02", tokens)
        second = task("B", "2026-09-01", "2026-09-02", tokens)
        result = self.analyze_pair(first, second)
        self.assertEqual({"room-1", "doe", "fixture-1"},
                         {value.resource_id for value in result.conflicts})

    def test_capacity_two_and_peak_three(self):
        nodes = [task(str(i), "2026-09-01", "2026-09-02",
                      [token("case-1", "C")]) for i in range(3)]
        definition = [ResourceDefinition("case-1", "case", "C", capacity=2)]
        two = analyze_projects([project("P", nodes[:2])], definition)
        self.assertFalse(two.conflicts)
        three = analyze_projects([project("P", nodes)], definition)
        self.assertEqual(3, three.conflicts[0].peak_usage)
        self.assertEqual(2, three.conflicts[0].capacity)
        self.assertEqual(3, len(three.conflicts[0].task_ids))

    def test_unchecked_identity_keeps_usage_but_does_not_flag_overlap(self):
        first = task("A", "2026-09-01", "2026-09-03",
                     [token("room-1", "Room 1", "room")])
        second = task("B", "2026-09-02", "2026-09-04",
                      [token("room-1", "Room 1", "room")])
        definition = [ResourceDefinition(
            "room-1", "room", "Room 1", conflict_enabled=False)]
        result = self.analyze_pair(first, second, definition)
        self.assertEqual(2, len(result.assignments))
        self.assertEqual([], result.conflicts)

    def test_work_calendar_parent_dateless_and_phase_exclusions(self):
        parent = task("Parent", "2026-09-04", "2026-09-08",
                      [token("case-1", "C")])
        child = task("Child", "2026-09-04", "2026-09-08",
                     [token("case-1", "C")], parent=parent)
        parent.children = [child]
        other = task("Other", "2026-09-04", "2026-09-08", [
            token("case-1", "C"), token("phase-1", "Design", "phase")])
        dateless = task("Inbox", None, None, [token("case-1", "C")])
        result = analyze_projects([
            project("P", [parent, other, dateless], holidays=["2026-09-07"])
        ])
        self.assertEqual(3, len(result.assignments))
        self.assertEqual(2, result.conflicts[0].overlap_workdays)
        self.assertEqual({parent.id, other.id},
                         set(result.conflicts[0].task_ids))
        self.assertEqual(2, result.conflicts[0].peak_usage)
        self.assertNotIn("phase-1", {a.resource_id for a in result.assignments})

    def test_projects_are_analyzed_together_and_fixed_dates_unchanged(self):
        first = task("Duplicate", "2026-09-01", "2026-09-03",
                     [token("case-1", "C")])
        second = task("Duplicate", "2026-09-02", "2026-09-04",
                      [token("case-1", "C")])
        second.start_rule = {"mode": "fixed", "date": second.start_date}
        before = (second.start_date, second.end_date, dict(second.start_rule))
        result = analyze_projects([project("P1", [first]), project("P2", [second])])
        self.assertEqual(1, len(result.conflicts))
        self.assertEqual({"P1", "P2"}, set(result.conflicts[0].project_ids))
        self.assertEqual(before, (second.start_date, second.end_date, second.start_rule))

    def test_acceptance_persists_and_changed_dates_invalidate_it(self):
        first = task("A", "2026-09-01", "2026-09-03", [token("case-1", "C")])
        second = task("B", "2026-09-02", "2026-09-04", [token("case-1", "C")])
        initial = self.analyze_pair(first, second)
        resolution = accept_conflict(initial.conflicts[0], "Approved shared setup")
        accepted = self.analyze_pair(first, second, resolutions=[resolution])
        self.assertEqual("accepted", accepted.conflicts[0].resolution_state)
        second.end_date = "2026-09-03"
        changed = self.analyze_pair(first, second, resolutions=[resolution])
        self.assertEqual("unresolved", changed.conflicts[0].resolution_state)
        self.assertNotEqual(initial.conflicts[0].assignment_signature,
                            changed.conflicts[0].assignment_signature)

    def test_next_available_search_only_returns_preview(self):
        first = task("A", "2026-09-01", "2026-09-03", [token("case-1", "C")])
        second = task("B", "2026-09-02", "2026-09-04", [token("case-1", "C")])
        result = self.analyze_pair(first, second)
        assignment = next(a for a in result.assignments if a.task_id == second.id)
        preview = next_available_start(assignment, result.assignments)
        self.assertEqual(("2026-09-04", "2026-09-08"), preview)
        self.assertEqual("2026-09-02", second.start_date)

    def test_deleted_task_removes_conflict(self):
        first = task("A", "2026-09-01", "2026-09-03", [token("case-1", "C")])
        second = task("B", "2026-09-02", "2026-09-04", [token("case-1", "C")])
        self.assertTrue(self.analyze_pair(first, second).conflicts)
        self.assertFalse(analyze_projects([project("P", [first])]).conflicts)

    def test_document_definitions_and_accepted_resolutions_round_trip(self):
        first = task("A", "2026-09-01", "2026-09-03", [token("case-1", "C")])
        second = task("B", "2026-09-02", "2026-09-04", [token("case-1", "C")])
        before = self.analyze_pair(first, second)
        definitions = [ResourceDefinition(
            "case-1", "case", "Renamed C", capacity=1,
            home_location="Room 1", notes="calibrated").to_dict()]
        resolutions = [accept_conflict(before.conflicts[0], "Approved overlap")]
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "resources.vpmt")
            save_projects([project("P", [first, second])], path,
                          resource_definitions=definitions,
                          conflict_resolutions=resolutions)
            loaded = load_projects(path)
        self.assertEqual(definitions, loaded[0]["resource_definitions"])
        self.assertEqual(resolutions,
                         loaded[0]["resource_conflict_resolutions"])
        after = analyze_projects(loaded, loaded[0]["resource_definitions"],
                                 loaded[0]["resource_conflict_resolutions"])
        self.assertEqual("accepted", after.conflicts[0].resolution_state)

    def test_legacy_file_without_document_definitions_loads_unchanged(self):
        original = task("Legacy", "2026-09-01", "2026-09-02", [])
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "legacy.vpmt")
            save_projects([project("P", [original])], path)
            loaded = load_projects(path)
        self.assertEqual([], loaded[0]["resource_definitions"])
        self.assertEqual("2026-09-01", loaded[0]["roots"][0].start_date)

    def test_multi_project_duplicate_task_id_conflict_attachment(self):
        task_p1 = task("Task in P1", "2026-09-01", "2026-09-03", [token("case-1", "C")])
        task_p1.id = "shared-id-1"
        task_p2 = task("Task in P2", "2026-09-02", "2026-09-04", [token("case-1", "C")])
        task_p2.id = "shared-id-1"

        p1 = project("P1", [task_p1], name="Project 1")
        p2 = project("P2", [task_p2], name="Project 2")

        result = analyze_projects([p1, p2], attach=True)
        self.assertEqual(1, len(result.conflicts))
        self.assertTrue(task_p1.schedule_conflicts)
        self.assertTrue(task_p2.schedule_conflicts)
        self.assertEqual(1, len(task_p1.resource_conflict_details))
        self.assertEqual(1, len(task_p2.resource_conflict_details))
        self.assertIn("Task in P2", task_p1.schedule_conflicts[0])
        self.assertIn("Task in P1", task_p2.schedule_conflicts[0])

    def test_f04_cassette_and_legacy_article_alias(self):
        node = task("Task with cassette", "2026-09-01", "2026-09-03", [
            {"id": "cass", "kind": "cassette", "label": "C", "header": "Cassettes"}
        ])
        node.resources = {"article": "C"}

        p = project("P", [node])
        result = analyze_projects([p])
        self.assertEqual(1, len(result.assignments))
        self.assertEqual("cass", result.assignments[0].resource_id)
        self.assertEqual(1, len(node.resource_tokens()))
        self.assertEqual("cass", node.resource_tokens()[0]["id"])

        # Multiple tokens with distinct IDs are kept distinct
        node2 = task("Two cassettes", "2026-09-01", "2026-09-03", [
            {"id": "cass-1", "kind": "cassette", "label": "C", "header": "Cassettes"},
            {"id": "cass-2", "kind": "cassette", "label": "C", "header": "Cassettes"},
        ])
        p2 = project("P2", [node2])
        res2 = analyze_projects([p2])
        self.assertEqual(2, len(res2.assignments))
        self.assertEqual({"cass-1", "cass-2"}, {a.resource_id for a in res2.assignments})

    def test_excel_resource_sheets_reconcile_to_analyzer(self):
        first = task("A", "2026-09-01", "2026-09-03", [token("case-1", "C")])
        second = task("B", "2026-09-02", "2026-09-04", [token("case-1", "C")])
        projects = [project("P", [first, second])]
        expected = analyze_projects(projects)
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "resources.xlsx")
            export_projects(projects, path)
            workbook = load_workbook(path, data_only=True)
        self.assertIn("Resource Usage", workbook.sheetnames)
        self.assertIn("Resource Conflicts", workbook.sheetnames)
        self.assertEqual(len(expected.assignments),
                         workbook["Resource Usage"].max_row - 1)
        self.assertEqual(len(expected.conflicts),
                         workbook["Resource Conflicts"].max_row - 1)


if __name__ == "__main__":
    unittest.main()
